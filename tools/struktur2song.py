#!/usr/bin/env python3
"""
struktur2song.py - song.yaml aus der Struktur-Tabelle des Users (Excel) erzeugen

    tools/.venv/Scripts/python tools/struktur2song.py <Song> --neu     # songs/<Song>/quelle/struktur.xlsx anlegen
                                                                       # (Kopie von songs/struktur-vorlage.xlsx)
    tools/.venv/Scripts/python tools/struktur2song.py <Song>           # Tabelle -> song.yaml

Die Tabelle ist der schlanke Nachfolger des Excel-Songkalkulators: der User schneidet die Parts so, wie er sie
für die Show braucht, und trägt nur ein, was er selbst weiß - Titel, Interpret, Song-ID, BPM, StartTakt/StartBit
(wo das Start-MIDI im DAW-Projekt liegt) und pro Part "bis Takt" (Taktnummer im DAW, an der der Part endet),
optional Energie, Effektidee, Beschreibung, Akkorde, Tempowechsel. Millisekunden rechnet songgen.py.

Schreibt songs/<Song>/song.yaml - aber NUR, wenn es die Datei noch nicht gibt. Eine vorhandene song.yaml
gehört dem User und wird nie überschrieben: das Ergebnis landet dann als Vorschlag in song.vorschlag.yaml
(der Generator ignoriert sie). Die Vorlage wird ebenfalls nie überschrieben.
"""
import argparse
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import markers as mk  # noqa: E402
from songgen import SONGS_DIR, SONG_FILE, VERSIONS_DIR, SongError, build_timeline, find_song_dir  # noqa: E402

TABLE_FILE = "struktur.xlsx"
TEMPLATE = SONGS_DIR / "struktur-vorlage.xlsx"	# Vorlage im Repo; --neu kopiert sie (der User darf sie anpassen)
PROPOSAL_FILE = "song.vorschlag.yaml"	# Ergebnis, wenn es schon eine song.yaml gibt (die wird nie überschrieben)
SHEET = "Struktur"
HEAD = [	# (Zeile, Beschriftung, Schlüssel, Erklärung)
	(1, "Titel", "name", ""),
	(2, "Interpret", "artist", ""),
	(3, "Song-ID", "id", "MIDI CC#0, 1..127 - freie Nummer oder die des alten Songs, den diese Fassung ersetzt"),
	(4, "BPM", "bpm", ""),
	(5, "StartTakt", "starttakt", "Taktnummer im DAW-Projekt, in der das Start-MIDI liegt"),
	(6, "StartBit", "startbit", "wie weit das MIDI nach dem Taktanfang kommt, als Bruchteil des Takts: 0,125 = 1/8 · 0,25 = 1/4 · 0,375 = 3/8"),
	(7, "Audio", "audio", "optional: Dateiname der MP3 in quelle/ (leer = die einzige MP3 im Ordner)"),
]
COLS = [	# (Überschrift, Schlüssel, Breite)
	("Part", "name", 30), ("bis Takt", "bis", 10), ("Takte", None, 8), ("Energie 0-5", "energy", 12),
	("Effektidee", "idea", 40), ("Beschreibung", "description", 44), ("Akkorde", "chords", 22), ("BPM", "bpm", 8),
]
HEADER_ROW = 9
FIRST_ROW = HEADER_ROW + 1
ROWS = 80		# so viele Part-Zeilen bekommt die Vorlage (mit Formel in "Takte")
NOTES = [
	"Parts so schneiden, wie du sie für die Show brauchst. 'bis Takt' = Taktnummer im DAW, an der der Part endet",
	"(= Anfang des nächsten). Halbe Takte als Komma-Zahl: 2269,5. Der erste Part beginnt am StartTakt.",
	"'Takte' rechnet nur zur Kontrolle. Kein Schluss-Black eintragen: 10 s Schwarz hängt der Generator an.",
	"BPM in der Part-Zeile nur bei Tempowechsel. Akkorde (z. B. Am F C G) nur für den Bund-Marker-Vorschlag.",
]

NOTE_IDX = {"C": 0, "C#": 1, "DB": 1, "D": 2, "D#": 3, "EB": 3, "E": 4, "F": 5, "F#": 6, "GB": 6,
			"G": 7, "G#": 8, "AB": 8, "A": 9, "A#": 10, "BB": 10, "B": 10, "H": 11}
NOTE_NAME = ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "H"]


#==================================================================
#=========== Vorlage ==============================================
#==================================================================

def write_template(path):
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = SHEET
	bold, grey = Font(bold=True), Font(color="808080", italic=True)
	fill = PatternFill("solid", fgColor="DDEBF7")
	for row, label, _key, hint in HEAD:
		ws.cell(row, 1, label).font = bold
		ws.cell(row, 2).fill = fill
		ws.cell(row, 3, hint).font = grey
	for c, (title, _key, width) in enumerate(COLS, 1):
		cell = ws.cell(HEADER_ROW, c, title)
		cell.font = bold
		cell.fill = PatternFill("solid", fgColor="BDD7EE")
		ws.column_dimensions[get_column_letter(c)].width = width
	for r in range(FIRST_ROW, FIRST_ROW + ROWS):
		prev = "$B$5" if r == FIRST_ROW else f"B{r - 1}"
		ws.cell(r, 3, f'=IF(B{r}="","",B{r}-{prev})').font = grey
		for c in (5, 6):
			ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
	ws.cell(FIRST_ROW, 1, "pause")
	for i, text in enumerate(NOTES):
		ws.cell(FIRST_ROW + i, len(COLS) + 2, text).font = grey
	ws.freeze_panes = ws.cell(FIRST_ROW, 2)
	path.parent.mkdir(parents=True, exist_ok=True)
	wb.save(path)


#==================================================================
#=========== Tabelle lesen ========================================
#==================================================================

def number(v, what, allow_fraction=False):
	if isinstance(v, bool) or v is None or v == "":
		raise SongError(f"{what} fehlt")
	if isinstance(v, (int, float)):
		return float(v)
	s = str(v).strip().replace(",", ".")
	if allow_fraction and "/" in s:
		n, d = s.split("/")
		return float(n) / float(d)
	try:
		return float(s)
	except ValueError:
		raise SongError(f"{what}: '{v}' ist keine Zahl")


def text(v):
	return "" if v is None else " ".join(str(v).split())


def read_table(path):
	wb = openpyxl.load_workbook(path, data_only=True)
	ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
	head = {key: ws.cell(row, 2).value for row, _label, key, _hint in HEAD}
	song = {"name": text(head["name"]), "artist": text(head["artist"]), "audio": text(head["audio"])}
	if not song["name"]:
		raise SongError("Titel fehlt")
	song["id"] = int(number(head["id"], "Song-ID"))
	song["bpm"] = number(head["bpm"], "BPM")
	start = number(head["starttakt"], "StartTakt")
	startbit = number(head["startbit"] if head["startbit"] not in (None, "") else 0, "StartBit", allow_fraction=True)
	if start != int(start) or not 0 <= startbit < 1:
		raise SongError("StartTakt muss eine ganze Taktnummer sein, StartBit der Bruchteil dahinter (0 bis unter 1)")
	song["starttakt"], song["startbit"] = int(start), startbit

	keys = {title: key for title, key, _w in COLS}
	cols = {keys[text(c.value)]: c.column for c in ws[HEADER_ROW] if text(c.value) in keys and keys[text(c.value)]}
	for need in ("name", "bis"):
		if need not in cols:
			raise SongError(f"Spalte '{'Part' if need == 'name' else 'bis Takt'}' fehlt in Zeile {HEADER_ROW}")
	parts, prev, seen = [], float(start), {}
	for r in range(FIRST_ROW, ws.max_row + 1):
		get = lambda k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
		name, bis = text(get("name")), get("bis")
		if not name and bis in (None, ""):
			continue
		if not name:
			raise SongError(f"Zeile {r}: Partname fehlt")
		if bis in (None, ""):
			if not parts and name == "pause":
				continue	# unausgefüllte Vorlage
			raise SongError(f"Zeile {r} '{name}': 'bis Takt' fehlt (bei Formeln: Datei einmal in Excel speichern)")
		bis = number(bis, f"Zeile {r} '{name}': bis Takt")
		if bis <= prev:
			raise SongError(f"Zeile {r} '{name}': bis Takt {fmt(bis)} liegt nicht nach {fmt(prev)}")
		seen[name.lower()] = seen.get(name.lower(), 0) + 1
		p = {"name": name if seen[name.lower()] == 1 else f"{name} ({seen[name.lower()]})", "von": prev, "bis": bis,
			 "bars": round(bis - prev, 4), "idea": text(get("idea")), "description": text(get("description")), "chords": text(get("chords"))}
		if get("energy") not in (None, ""):
			p["energy"] = int(number(get("energy"), f"Zeile {r} '{name}': Energie"))
			if not 0 <= p["energy"] <= 5:
				raise SongError(f"Zeile {r} '{name}': Energie {p['energy']} liegt nicht in 0..5")
		if get("bpm") not in (None, ""):
			p["bpm"] = number(get("bpm"), f"Zeile {r} '{name}': BPM")
		parts.append(p)
		prev = bis
	if not parts:
		raise SongError("keine Parts eingetragen")
	song["parts"] = parts
	return song


#==================================================================
#=========== song.yaml schreiben ==================================
#==================================================================

def fmt(x):
	return f"{x:.4f}".rstrip("0").rstrip(".")


def ystr(s):
	return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def offset_lines(song):
	"""StartBit (Bruchteil des Takts) -> midi_offset als Notenwert; krumme Werte als midi_offset_ms."""
	sixteenths = song["startbit"] * 16
	if abs(sixteenths - round(sixteenths)) > 1e-6:
		bpm = song["parts"][0].get("bpm", song["bpm"])
		return [f"midi_offset_ms: {round(song['startbit'] * 240000.0 / bpm)}        # StartBit {fmt(song['startbit'])} ist kein Notenwert"]
	n, d = round(sixteenths), 16
	while n and n % 2 == 0:
		n, d = n // 2, d // 2
	return [f"midi_offset: {f'{n}/{d}' if n else 0}"]


def length_lines(bars, bpb=4):
	whole = int(bars + 1e-9)
	beats = round((bars - whole) * bpb, 3)
	return ([f"    bars: {whole}"] if whole or not beats else []) + ([f"    beats: {fmt(beats)}"] if beats else [])


def parse_chord(tok):
	"""'C#m7/G#' -> (1, 'min'); None, wenn kein Akkord."""
	m = re.match(r"\s*([A-Ha-h])([#b]?)(.*)$", tok)
	root = NOTE_IDX.get((m.group(1).upper() + m.group(2).upper()).replace("HB", "B")) if m else None
	if root is None:
		return None
	return (root, "min" if re.match(r"(m(?!aj)|min|-|dim)", m.group(3).split("/")[0]) else "maj")


def marker_lines(song, song_file):
	"""Bund-Marker: nie etwas Vorhandenes ändern (wie sheet2song.py)."""
	counts = {}
	for p in song["parts"]:
		for tok in re.split(r"[\s|,]+", p["chords"]):
			ch = parse_chord(tok) if tok else None
			if ch:
				counts[ch] = counts.get(ch, 0) + p["bars"]
	chords_text = ", ".join(NOTE_NAME[c[0]] + ("m" if c[1] == "min" else "") for c, _n in sorted(counts.items(), key=lambda kv: -kv[1]))
	if song["id"] in mk.handwritten_ids():
		return [f"# Bund-Marker: von Hand in src/markerLEDs.cpp gesetzt (case {song['id']}), werden dort gepflegt."] \
			+ ([f"# Akkorde im Song: {chords_text}"] if chords_text else []), "von Hand in markerLEDs.cpp - kein Vorschlag"
	if song_file.exists():
		m = re.search(r"^(?:#[^\n]*\n)*markers:.*", song_file.read_text(encoding="utf-8"), re.S | re.M)
		if m:
			return [m.group(0).rstrip("\n")], "aus der vorhandenen song.yaml übernommen (unverändert)"
	if not counts:
		return ["# Bund-Marker: noch keine. Akkorde in der Tabelle eintragen (Spalte Akkorde) oder hier markers: ergänzen."], \
			"keine (Spalte Akkorde ist leer)"
	names, notes = mk.propose(counts)
	return mk.yaml_block(names, notes, chords_text), "VORSCHLAG: " + ", ".join(f"{n} ({mk.FRET[n]}. Bund)" for n in names)


def find_mp3(song, song_dir):
	if song["audio"]:
		for c in (song_dir / "quelle" / song["audio"], song_dir / song["audio"]):
			if c.is_file():
				return c
		raise SongError(f"Audio '{song['audio']}' nicht gefunden (gesucht in quelle/ und im Song-Ordner)")
	found = [p for p in song_dir.rglob("*.mp3") if VERSIONS_DIR not in p.parts]
	return found[0] if len(found) == 1 else None


def render(song, song_dir, table, is_proposal):
	song_file = song_dir / SONG_FILE
	audio = find_mp3(song, song_dir)
	lines = []
	if is_proposal:
		lines += [f"# VORSCHLAG - {SONG_FILE} existiert schon und wurde NICHT angefasst. Nur zum Vergleichen,",
				  "# der Generator beachtet diese Datei nicht. Übernimm von Hand, was du brauchst.", ""]
	lines += [
		f"# Semantische Songbeschreibung, erzeugt von tools/struktur2song.py aus quelle/{table.name}",
		f"# (StartTakt {song['starttakt']}, StartBit {fmt(song['startbit'])}; Parts bis Takt {fmt(song['parts'][-1]['bis'])}).",
		"# Die Technik (Szenen/Farben) leitet Claude in show.yaml ab. Diese Datei hier gehört dir: kein Werkzeug",
		"# und kein Claude überschreibt sie. Pro Part kannst du ergänzen:",
		"#   energy: 0-5, description, idea, mood, instruments, solo, lyrics  -> Einschätzung, daraus wird die Show abgeleitet",
		"#   scene: SCENE_..., scheme: SCHEME_..., fx: \"prog...\"  -> feste Vorgabe, hat immer Vorrang vor show.yaml",
		"",
		f"id: {song['id']}", f"name: {ystr(song['name'])}", f"artist: {ystr(song['artist'])}", f"bpm: {fmt(song['bpm'])}",
		"beats_per_bar: 4",
	] + offset_lines(song)
	if audio:
		lines += [f"audio: {audio.relative_to(song_dir).as_posix()}",
				  "# audio_beat1_ms: 0          # wo Takt 1 in der MP3 liegt - songanalyze.py schätzt es, wenn die Zeile fehlt"]
	lines += ["", "sections:"]
	for p in song["parts"]:
		lines.append(f"  - name: {ystr(p['name'])}")
		lines += length_lines(p["bars"])
		if "bpm" in p and abs(p["bpm"] - song["bpm"]) > 0.001:
			lines.append(f"    bpm: {fmt(p['bpm'])}")
		if "energy" in p:
			lines.append(f"    energy: {p['energy']}")
		for k in ("idea", "description", "chords"):
			if p[k]:
				lines.append(f"    {k}: {ystr(p[k])}")
		lines.append("")
	block, msg = marker_lines(song, song_file)
	return "\n".join(lines + block) + "\n", msg


def timeline_table(song):
	"""Kontrollausgabe wie im alten Kalkulator: Takte und Millisekunden ab dem Start-MIDI (so rechnet songgen.py)."""
	sixteenths = song["startbit"] * 16
	gen = {"bpm": song["bpm"], "sections": [dict({"name": p["name"], "bars": p["bars"]}, **({"bpm": p["bpm"]} if "bpm" in p else {}))
											 for p in song["parts"]]}
	if abs(sixteenths - round(sixteenths)) > 1e-6:
		gen["midi_offset_ms"] = song["startbit"] * 240000.0 / song["parts"][0].get("bpm", song["bpm"])
	else:
		gen["midi_offset"] = f"{round(sixteenths)}/16" if round(sixteenths) else 0
	timeline, _end = build_timeline(gen)
	print(f"\n  {'Part':<32} {'von':>8} {'bis':>8} {'Takte':>7} {'Start ms':>9} {'Dauer ms':>9}")
	for p, t in zip(song["parts"], timeline):
		von = p["von"] + (song["startbit"] if p is song["parts"][0] else 0)
		print(f"  {p['name'][:32]:<32} {fmt(von):>8} {fmt(p['bis']):>8} {fmt(p['bis'] - von):>7} {t['start']:>9} {t['dur']:>9}")
	total = sum(p["bars"] for p in song["parts"])
	print(f"  {len(song['parts'])} Parts, {fmt(total)} Takte, {timeline[len(song['parts']) - 1]['start'] + timeline[len(song['parts']) - 1]['dur']} ms "
		  f"(+ 10 s Schwarz am Ende)")


def main():
	sys.stdout.reconfigure(encoding="utf-8")
	ap = argparse.ArgumentParser(description="song.yaml aus der Struktur-Tabelle (Excel) des Users erzeugen")
	ap.add_argument("song", help="Ordnername unter songs/ (mit --neu: genau so wird er angelegt, z. B. Vogue_v1)")
	ap.add_argument("--neu", action="store_true", help=f"leere Vorlage quelle/{TABLE_FILE} anlegen")
	args = ap.parse_args()
	try:
		if args.neu:
			path = SONGS_DIR / args.song / "quelle" / TABLE_FILE
			if path.exists():
				raise SongError(f"{path.relative_to(SONGS_DIR.parent)} gibt es schon - sie wird nicht überschrieben")
			if TEMPLATE.is_file():
				path.parent.mkdir(parents=True, exist_ok=True)
				shutil.copyfile(TEMPLATE, path)
			else:
				write_template(path)
			print(f"Vorlage angelegt: {path.relative_to(SONGS_DIR.parent)}\nAusfüllen, speichern, dann: struktur2song.py {args.song}")
			return
		song_dir = find_song_dir(args.song)
		table = next((p for p in (song_dir / "quelle" / TABLE_FILE, song_dir / TABLE_FILE) if p.is_file()), None)
		if table is None:
			raise SongError(f"songs/{song_dir.name}/quelle/{TABLE_FILE} fehlt - Vorlage anlegen mit --neu")
		song = read_table(table)
		is_proposal = (song_dir / SONG_FILE).exists()
		out = song_dir / (PROPOSAL_FILE if is_proposal else SONG_FILE)
		content, marker_msg = render(song, song_dir, table, is_proposal)
		print(f"#{song['id']} {song['name']} - {song['artist']}, {fmt(song['bpm'])} BPM, {offset_lines(song)[0].split('#')[0].strip()}")
		timeline_table(song)
		out.write_text(content, encoding="utf-8")
		print(f"\nBund-Marker: {marker_msg}")
		if is_proposal:
			print(f"{SONG_FILE} existiert schon und bleibt unangetastet -> Vorschlag: {out.relative_to(SONGS_DIR.parent)}")
		else:
			print(f"geschrieben: {out.relative_to(SONGS_DIR.parent)}")
	except SongError as e:
		raise SystemExit(f"FEHLER: {e}")


if __name__ == "__main__":
	main()
