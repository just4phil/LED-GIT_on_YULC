#!/usr/bin/env python3
"""
excel2song.py - einmalige Übernahme der alten Songs in die Ordnerstruktur songs/<Song>_v1/

    python tools/excel2song.py <Songkalkulator.xlsx>             # nur Bericht, schreibt nichts
    python tools/excel2song.py <Songkalkulator.xlsx> --write     # song.yaml + quelle/excel-kalkulation.csv anlegen
    python tools/excel2song.py <Songkalkulator.xlsx> --write Kids  # nur diese Songs (Funktionsname aus songs.cpp)

Quellen: das Blatt "Tabelle1 (2)" des Excel-Songkalkulators (BPM, StartBit, Partnamen, Takte) und die
handgeschriebenen Song-Funktionen in src/songs.cpp (Dauern in ms, live erprobt). Wo beide auseinanderlaufen,
gilt der Code: seine Dauern werden mit dem Tempo in Takte umgerechnet.

Schreibt songs/<Song>_v1/song.yaml NUR, wenn es die Datei noch nicht gibt - eine vorhandene song.yaml gehört
dem User und wird übersprungen (wie bei sheet2song.py). Braucht openpyxl.
"""
import argparse
import csv
import datetime
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CPP = ROOT / "src" / "songs.cpp"
SONGS_DIR = ROOT / "songs"
ARCHIVE_DIR = ROOT / "docs" / "songkalkulator"
SHEET = "Tabelle1 (2)"
SNAP = 8				# Raster für Taktlängen aus dem Code: 1/8 Takt
SNAP_TOL_MS = 35		# so weit darf eine Code-Dauer vom Raster bzw. von der Excel-Zeile abweichen (1/8 Takt >= 190 ms)
END_NAMES = ("black", "fini", "finito", "ende", "back to default")

# Funktion in songs.cpp -> (Ordner, Song-ID, Titel, Block im Excel, Hinweis)
SONGS = {
	"TakeOnMe": ("TakeOnMe_v1", 3, "Take On Me", "TAKE ON ME", ""),
	"DontStopTheMusic": ("DontStopTheMusic_v1", 4, "Don't Stop The Music", "Dont Stop the music", ""),
	"NoRoots": ("NoRoots_v1", 6, "No Roots", "No Roots", ""),
	"Firework": ("Firework_v1", 7, "Firework", "firework", ""),
	"DancingOnMyOwn": ("DancingOnMyOwn_v1", 8, "Dancing On My Own", "Dancing on my own",
					   "Intro/Trailer: INTROdancing() (#81) springt in diesen Song"),
	"ILoveIt": ("ILoveIt_v1", 9, "I Love It", "I LOVE IT", "Intro/Trailer: ILoveItTRAILER() (#80) springt in diesen Song"),
	"BloodyMary": ("BloodyMary_v1", 10, "Bloody Mary", "BloodyMary", ""),
	"Titanium": ("Titanium_v1", 11, "Titanium", "Titanium", ""),
	"SuchAshame": ("SuchAShame_v1", 12, "Such A Shame", "Such a Shame", ""),
	"InTheDark": ("InTheDark_v1", 13, "In The Dark", "In The Dark", ""),
	"Shivers": ("Shivers_v1", 14, "Shivers", "SHIVERS", ""),
	"Abcdefu": ("Abcdefu_v1", 15, "abcdefu", "abcdefu", ""),
	"enjoyTheSilence": ("EnjoyTheSilence_v1", 16, "Enjoy The Silence", "Enjoy the silence",
						"Intro/Trailer: enjoyTheSilenceINTRO() (#24) springt in diesen Song"),
	"apt": ("Apt_v1", 17, "APT.", "APT", ""),
	"Kids": ("Kids_v1", 20, "Kids", "KIDS", ""),
	"Tellittomyheart": ("TellItToMyHeart_v1", 21, "Tell It To My Heart", "TELL IT TO MY HEART", ""),
	"FridayImInLove": ("FridayImInLove_v1", 25, "Friday I'm In Love", "Friday im in Love", ""),
	"BeMine": ("BeMine_v1", 26, "Be Mine", "Be Mine", ""),
	"IWannaDanceWithSomebody": ("IWannaDanceWithSomebody_v1", 27, "I Wanna Dance With Somebody", "I Wanna dance with somebody", ""),
	"BillyJean": ("BillieJean_v1", 28, "Billie Jean", "Billie Jean", ""),
	"Maniac_Tminus1": ("Maniac_v1", 30, "Maniac", "Maniac", "Maniac_Tminus1() (#30) ist die genutzte Fassung, Maniac() (#29) ist veraltet"),
}
ARTISTS = {"enjoyTheSilence": "Depeche Mode"}	# wo der Lauftext nicht in der Song-Funktion steht
# nur die Excel-Kalkulation archivieren, song.yaml gibt es schon
ARCHIVE_ONLY = {"Physical_v1": "Physical (OHNE TRAILER)"}


#==================================================================
#=========== Excel ================================================
#==================================================================

def num(v):
	return isinstance(v, (int, float)) and not isinstance(v, bool)


def load_excel(path):
	ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
	rows = [list(r) + [None] * 12 for r in ws.iter_rows(values_only=True)]
	heads = [i for i, r in enumerate(rows) if r[1] == "BPM" and r[3] == "StartTakt"]
	blocks = {}
	for n, i in enumerate(heads):
		end = heads[n + 1] - 1 if n + 1 < len(heads) else len(rows)
		title = rows[i - 1]
		b = {"name": str(title[1]).strip(), "note": " ".join(str(c) for c in title[2:8] if c), "bpm": rows[i + 1][1],
			 "starttakt": rows[i + 1][3], "startbit": rows[i + 2][3], "parts": {}, "rows": rows[i - 1:end]}
		for p in rows[i + 5:end]:
			if num(p[0]) and num(p[7]):
				b["parts"][int(p[0])] = {"name": str(p[1] or "").strip(), "ms": p[2], "fx": p[3], "bpm": p[7], "takte": p[10]}
		blocks[b["name"]] = b
	return blocks


#==================================================================
#=========== songs.cpp ============================================
#==================================================================

NEXT_IDX = {"progFastBlingBling": 2, "progSternNeu": 2, "progPalette": 2, "progScrollText": 4, "progShowText": 5,
			"progWordArray": 5, "progShowLettersSpread": 2, "progBlinkText": 2}
TEXT_FX = ("progScrollText", "progShowText", "progWordArray", "progShowLettersSpread", "progBlinkText")


def strip_comments(src):
	"""Kommentare entfernen; den Kommentar hinter 'case N:' als Partnamen merken."""
	src = re.sub(r"/\*.*?\*/", lambda m: "\n" * m.group(0).count("\n"), src, flags=re.S)
	out = []
	for line in src.split("\n"):
		res, quoted, k = "", False, 0
		while k < len(line):
			c = line[k]
			if c == '"':
				quoted = not quoted
			if not quoted and line[k:k + 2] == "//":
				if re.match(r"\s*case \d+\s*:\s*$", res):
					res += " /*N:" + line[k + 2:].strip().replace("*/", "") + "*/"
				break
			res += c
			k += 1
		out.append(res)
	return "\n".join(out)


def split_args(s):
	args, depth, quoted, cur = [], 0, False, ""
	for c in s:
		if c == '"':
			quoted = not quoted
		if not quoted:
			depth += c in "(["
			depth -= c in ")]"
			if c == "," and depth == 0:
				args.append(cur.strip())
				cur = ""
				continue
		cur += c
	args.append(cur.strip())
	return args


def load_funcs():
	src = strip_comments(CPP.read_text(encoding="utf-8", errors="replace"))
	funcs = {}
	for m in re.finditer(r"^void (\w+)\(\)\s*\{(.*?)^\}", src, flags=re.S | re.M):
		pieces = re.split(r"\bcase (\d+)\s*:", m.group(2))
		cases = {}
		for k in range(1, len(pieces), 2):
			text = pieces[k + 1]
			name = re.match(r"\s*/\*N:(.*?)\*/", text)
			calls = []
			for c in re.finditer(r"\b(prog\w+|matrixMovieFX)\s*\((.*?)\)\s*;", text, flags=re.S):
				fx, a = c.group(1), split_args(" ".join(c.group(2).split()))
				di = 3 if fx == "progWordArray" else 1 if fx in TEXT_FX else 0
				ni = 2 if fx == "progStern" and len(a) == 4 else NEXT_IDX.get(fx, 1)
				try:
					dur, nxt = float(a[di]), int(a[ni])
				except (ValueError, IndexError):
					continue
				calls.append({"fx": fx, "dur": dur, "next": nxt, "else": bool(re.search(r"\belse\b", text[:c.start()])),
							  "raw": f"{fx}({', '.join(a)})"})
			cases[int(pieces[k])] = {"name": name.group(1) if name else "", "calls": calls,
									  "markers": sorted(set(re.findall(r"markerLED\d\s*=\s*\w+", text)))}
		funcs[m.group(1)] = {"cases": cases, "text": m.group(2)}
	return funcs


def guitar_chain(cases):
	"""Ablauf auf Gitarre/Bass: bei mehreren Aufrufen der else-Zweig (nicht die Matrix), kein Lauftext."""
	chain, c, seen = [], 0, set()
	while c in cases and c not in seen and cases[c]["calls"]:
		seen.add(c)
		calls = cases[c]["calls"]
		pick = ([x for x in calls if x["else"]] or [x for x in calls if x["fx"] not in TEXT_FX] or calls)[0]
		chain.append({"case": c, "name": clean_name(cases[c]["name"]), "call": pick, "markers": cases[c]["markers"],
					  "other": [x["raw"] for x in calls if x is not pick and x["fx"] != pick["fx"]]})
		c = pick["next"]
	return chain


def clean_name(s):
	"""'85\\tBLACK\\t10000' -> 'BLACK', 'so far away, 15735' -> 'so far away', 'verse 2b, 11800 -> 19670' -> 'verse 2b'."""
	tokens = [t.strip() for t in s.split("\t") if t.strip() and not re.fullmatch(r"[\d.]+", t.strip())]
	s = tokens[0] if tokens else ""
	s = re.sub(r"^\d+\s+(?=\D)", "", s)
	return re.sub(r"[\s,:]+\d{3,}.*$", "", s).strip(" ,:-")


#==================================================================
#=========== Struktur ableiten ====================================
#==================================================================

def snap(bars):
	return round(bars * SNAP) / SNAP


def derive(fn, funcs, xl):
	folder, song_id, title, xname, hint = SONGS[fn]
	b = xl[xname]
	chain = guitar_chain(funcs[fn]["cases"])
	notes = []
	end_black = black_case = None
	if chain and chain[-1]["call"]["fx"] == "progBlack":
		black_case = chain[-1]["case"]
		end_black = int(chain.pop()["call"]["dur"])
	offset = float(b["startbit"]) % 1
	bpm = float(b["bpm"])
	sections = []
	t_code = 0.0
	matched = 0
	i = 0

	def sec(name, bars, src, parts, flag=""):
		return {"name": name, "bars": bars, "bpm": bpm, "src": src, "parts": parts, "flag": flag, "t_end": None}

	def excel_rows(lo, hi, total, count):
		"""Excel-Zeilen der cases lo..hi-1, wenn ihre Summe der Code-Dauer entspricht."""
		rows = [(c, p) for c, p in sorted(b["parts"].items()) if lo <= c < hi]
		ok = rows and all(num(p["ms"]) and num(p["takte"]) and p["takte"] > 0 for _, p in rows)
		return rows if ok and abs(sum(p["ms"] for _, p in rows) - total) <= SNAP_TOL_MS * max(count, len(rows)) else None

	while i < len(chain):
		e = chain[i]
		off = offset if i == 0 else 0.0
		used = None
		# 1. Excel und Code stimmen überein: ein Code-Part = eine oder mehrere Excel-Zeilen; sonst 2. die Code-Dauer liegt
		#    im Taktraster; sonst 3. mehrere Code-Parts zusammen = Excel-Zeilen bzw. zusammen im Taktraster
		for j in [i] + [None] + list(range(i + 1, min(i + 4, len(chain)))):
			if j is None:
				xp = b["parts"].get(e["case"])
				bpm = float(xp["bpm"]) if xp else bpm
				bar = 240000.0 / bpm
				bars = e["call"]["dur"] / bar + off
				if abs(bars - snap(bars)) * bar <= SNAP_TOL_MS:
					flag = ""
					same = xp and (not e["name"] or e["name"].lower() == xp["name"].lower())	# sonst ist die Nummerierung im Excel verrutscht
					if same and num(xp["takte"]) and xp["takte"] > 0 and abs(xp["takte"] + off - snap(bars)) > 0.01:
						flag = f"Excel: {fmt(xp['takte'] + off)} Takte, der alte Code ist anders"
						notes.append(f"case {e['case']} '{e['name'] or xp['name']}': Code {fmt(snap(bars))} Takte, Excel {fmt(xp['takte'] + off)}")
					sections.append(sec(e["name"] or (xp["name"] if xp else f"part {e['case']}"), snap(bars), "code", [e], flag))
					used = 1
					break
				continue
			group = chain[i:j + 1]
			total = sum(g["call"]["dur"] for g in group)
			upper = chain[j + 1]["case"] if j + 1 < len(chain) else (black_case or 10 ** 6)
			rows = excel_rows(e["case"], upper, total, len(group))
			if rows:
				split = len(rows) != len(group) or len(group) > 1
				for n, (c, p) in enumerate(rows):
					bpm = float(p["bpm"])
					name = (e["name"] if not split else "") or p["name"]
					flag = ""
					if split and n == 0:
						flag = (f"der alte Code teilt diese {len(rows)} Parts anders auf (Summe identisch)" if len(group) > 1
								else "im alten Code ein einziger Part zusammen mit " + ("dem folgenden" if len(rows) == 2 else f"den folgenden {len(rows) - 1}"))
						notes.append(f"case {e['case']}..{group[-1]['case']}: Code {len(group)} Part(s), Excel {len(rows)} - Excel-Teilung übernommen")
					sections.append(sec(name, round(p["takte"] + (off if n == 0 else 0), 4), "excel", group if n == 0 else [], flag))
				matched += len(group)
				used = len(group)
				break
			if j > i:
				bar = 240000.0 / bpm
				bars = total / bar + off
				if abs(bars - snap(bars)) * bar <= SNAP_TOL_MS * len(group):
					names = [g["name"] for g in group if g["name"]]
					sections.append(sec(" + ".join(dict.fromkeys(names)) or f"part {e['case']}", snap(bars), "code", group,
										"im alten Code mehrere Parts mit krummer Länge, hier zusammengefasst"))
					notes.append(f"case {e['case']}..{group[-1]['case']}: zusammengefasst zu {fmt(snap(bars))} Takten")
					used = len(group)
					break
		if not used:
			bars = e["call"]["dur"] * bpm / 240000.0 + off
			sections.append(sec(e["name"] or f"part {e['case']}", round(bars, 3), "code", [e], "krumme Länge im alten Code - bitte prüfen"))
			notes.append(f"case {e['case']} '{e['name']}': {e['call']['dur']:.0f} ms = {bars:.3f} Takte (krumm)")
			used = 1
		t_code += sum(g["call"]["dur"] for g in chain[i:i + used])
		sections[-1]["t_end"] = t_code
		i += used

	# Parts, die es nur im Excel gibt (der alte Code hört vorher auf) - nur wenn Excel und Code sonst zusammenpassen
	if chain and matched >= 0.8 * len(chain):
		for c, p in sorted(b["parts"].items()):
			if c > chain[-1]["case"] and c != black_case and num(p["takte"]) and p["takte"] > 0 and not p["name"].lower().startswith(END_NAMES) and p["name"]:
				bpm = float(p["bpm"])
				sections.append({"name": p["name"], "bars": round(p["takte"], 4), "bpm": bpm, "src": "excel", "parts": [], "t_end": None,
								 "flag": "nur im Excel - der alte Code geht hier schon in Schwarz"})
				notes.append(f"case {c} '{p['name']}' ({fmt(p['takte'])} Takte) steht nur im Excel")
	elif chain:
		notes.append(f"Excel und Code passen nur bei {matched} von {len(chain)} Parts zusammen - Struktur aus dem Code")

	# eindeutige Namen
	seen = {}
	for s in sections:
		s["name"] = s["name"] or "part"
		seen[s["name"]] = seen.get(s["name"], 0) + 1
		if seen[s["name"]] > 1:
			s["name"] += f" ({seen[s['name']]})"

	# Gegenprobe: Grenzen so rechnen wie songgen.py und mit den Summen des alten Codes vergleichen
	off_ms = offset * 240000.0 / sections[0]["bpm"]
	t, drift = -off_ms, 0.0
	for s in sections:
		t += s["bars"] * 240000.0 / s["bpm"]
		if s["t_end"] is not None:
			drift = max(drift, abs(round(t) - s["t_end"]))

	scroll = re.search(r'progScrollText\("([^"]+?) by ([^"]+)"', funcs[fn]["text"])
	return {"fn": fn, "folder": folder, "id": song_id, "name": title, "artist": scroll.group(2).strip() if scroll else ARTISTS.get(fn, ""),
			"bpm": b["bpm"], "offset": offset, "block": b, "sections": sections, "end_black": end_black, "notes": notes,
			"drift": drift, "hint": hint, "code_ms": t_code, "inline_markers": any(e["markers"] for e in chain)}


#==================================================================
#=========== Ausgabe ==============================================
#==================================================================

def fmt(x):
	return f"{x:.4f}".rstrip("0").rstrip(".")


def offset_note(offset):
	if not offset:
		return "0"
	n = round(offset * 16)
	d = 16
	while n % 2 == 0:
		n //= 2
		d //= 2
	return f"{n}/{d}"


def q(s):
	return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def length_lines(sec, bpb=4):
	whole = int(sec["bars"] + 1e-9)
	beats = round((sec["bars"] - whole) * bpb, 3)
	lines = []
	if whole or not beats:
		lines.append(f"    bars: {whole}")
	if beats:
		lines.append(f"    beats: {fmt(beats)}")
	return lines


def render(song):
	b = song["block"]
	out = [
		f"# Grundstruktur aus dem Excel-Songkalkulator und {song['fn']}() in src/songs.cpp (tools/excel2song.py, "
		f"{datetime.date.today():%d.%m.%Y}).",
		"# Die Längen entsprechen dem alten, live erprobten Code. Zeilen mit '# !' bitte prüfen.",
		"# '# bisher:' zeigt den Effekt des alten Codes auf der Gitarre - nur zur Orientierung, der Generator liest das nicht.",
		"# Die Technik (Szenen/Farben) leitet Claude in show.yaml ab. Diese Datei hier gehört dir: kein Werkzeug",
		"# und kein Claude überschreibt sie. Pro Part kannst du ergänzen:",
		"#   energy: 0-5, description, mood, instruments, solo, lyrics  -> Einschätzung, daraus wird die Show abgeleitet",
		"#   scene: SCENE_..., scheme: SCHEME_..., fx: \"prog...\"  -> feste Vorgabe, hat immer Vorrang vor show.yaml",
		f"# Excel: StartTakt {b['starttakt']}, StartBit {fmt(float(b['startbit']))}" + (f" - {b['note']}" if b["note"] else ""),
	]
	if song["hint"]:
		out.append(f"# {song['hint']}")
	out += ["", f"id: {song['id']}", f"name: {q(song['name'])}", f"artist: {q(song['artist'])}", f"bpm: {fmt(float(song['bpm']))}",
			"beats_per_bar: 4", f"midi_offset: {offset_note(song['offset'])}"]
	if song["end_black"] not in (None, 10000):
		out.append(f"end_black_ms: {song['end_black']}")
	out += ["", "sections:"]
	for s in song["sections"]:
		out.append(f"  - name: {q(s['name'])}")
		out += length_lines(s)
		if abs(s["bpm"] - float(song["bpm"])) > 0.01:
			out.append(f"    bpm: {fmt(s['bpm'])}")
		for e in s["parts"]:
			out.append(f"    # bisher: {e['call']['raw']}" + (f"   (Matrix: {'; '.join(e['other'])})" if e["other"] else ""))
			if e["markers"]:
				out.append(f"    # ! der alte Code setzt hier Marker: {'; '.join(e['markers'])}")
		if s["flag"]:
			out.append(f"    # ! {s['flag']}")
		out.append("")
	out.append(f"# Bund-Marker: im alten Code in src/markerLEDs.cpp (setMarkerLEDs, Song {song['id']})"
			   + (" und direkt in einzelnen Parts (siehe '# !' oben)." if song["inline_markers"] else "."))
	out.append("# Beim Umstellen auf den Generator müssen sie 1:1 übernommen werden (markers: in dieser Datei).")
	return "\n".join(out) + "\n"


def write_csv(path, rows):
	path.parent.mkdir(parents=True, exist_ok=True)
	with open(path, "w", newline="", encoding="utf-8-sig") as f:
		w = csv.writer(f, delimiter=";")
		for r in rows:
			r = list(r)
			while r and r[-1] is None:
				r.pop()
			w.writerow(["" if c is None else fmt(c) if isinstance(c, float) else c for c in r])


def main():
	sys.stdout.reconfigure(encoding="utf-8")
	ap = argparse.ArgumentParser(description="Alte Songs (Excel-Songkalkulator + songs.cpp) nach songs/<Song>_v1/ übernehmen")
	ap.add_argument("xlsx")
	ap.add_argument("songs", nargs="*", help="Funktionsnamen aus songs.cpp (Standard: alle)")
	ap.add_argument("--write", action="store_true", help="Dateien anlegen (sonst nur Bericht)")
	args = ap.parse_args()

	xl, funcs = load_excel(args.xlsx), load_funcs()
	for fn in SONGS:
		if args.songs and fn not in args.songs:
			continue
		song = derive(fn, funcs, xl)
		target = SONGS_DIR / song["folder"] / "song.yaml"
		src = {k: sum(1 for s in song["sections"] if s["src"] == k) for k in ("excel", "code")}
		bars = sum(s["bars"] for s in song["sections"])
		print(f"\n#{song['id']:<3} {song['folder']:<28} {fmt(float(song['bpm']))} BPM  midi_offset {offset_note(song['offset'])}  "
			  f"{len(song['sections'])} Parts, {fmt(bars)} Takte  (Excel {src['excel']}, Code {src['code']})  "
			  f"Abweichung zum alten Code max. {song['drift']:.0f} ms  artist={song['artist']!r}")
		for n in song["notes"]:
			print(f"      ! {n}")
		if not args.write:
			continue
		if target.exists():
			print("      song.yaml vorhanden -> übersprungen")
		else:
			target.parent.mkdir(parents=True, exist_ok=True)
			target.write_text(render(song), encoding="utf-8")
			print(f"      geschrieben: {target.relative_to(ROOT)}")
		write_csv(target.parent / "quelle" / "excel-kalkulation.csv", song["block"]["rows"])

	if args.write and not args.songs:
		for folder, xname in ARCHIVE_ONLY.items():
			write_csv(SONGS_DIR / folder / "quelle" / "excel-kalkulation.csv", xl[xname]["rows"])
		wb = openpyxl.load_workbook(args.xlsx, data_only=True)
		for ws in wb:
			write_csv(ARCHIVE_DIR / (re.sub(r"[^\w-]+", "_", ws.title).strip("_") + ".csv"), ws.iter_rows(values_only=True))
		print(f"\nKomplette Tabelle archiviert: {ARCHIVE_DIR.relative_to(ROOT)}/")


if __name__ == "__main__":
	main()
